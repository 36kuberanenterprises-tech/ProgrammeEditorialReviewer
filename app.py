import io
import os
import re
import json
import difflib
from pathlib import Path
from typing import List, Dict, Any, Optional

import httpx
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from fastapi.requests import Request
from docx import Document
from docx.shared import RGBColor
from pypdf import PdfReader

BASE_DIR = Path(__file__).resolve().parent
app = FastAPI(title="Programme Editorial Reviewer")
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))

SYSTEM_PROMPT = r"""
You are a senior Indian NGO, CSR and development sector editor. This is not basic grammar correction. Review the supplied programme text as an institutional editor.

IMPORTANT CONTEXT RULE:
The paragraphs are part of one document, chapter or subject flow. Do not treat each paragraph as isolated. Use the supplied previous and following context to preserve continuity, sequence and links between paragraphs. Do not remove an explanation from one paragraph if it is needed to introduce, support or connect with the next paragraph. Judge repetition across the surrounding section, not only inside one paragraph.

Rules:
1. Preserve every verified fact, figure, name, location, date, income figure, funding amount and programme term exactly as supplied.
2. Place names are protected details. Preserve relevant names of states, districts, taluks, blocks, towns, villages, Gram Panchayats, cities, programme locations, training centres, markets and other geographic references across India and in any other country appearing in the source. Do not remove or generalise a place name when it represents geography, beneficiary background, programme coverage, field evidence, implementation location or institutional context. Reduce a repeated place name only when the location is already completely clear and the repetition serves no purpose.
3. Never invent achievements, referrals, customer growth, confidence, demand, family support, impact, partnerships or outcomes.
4. If a claim needs verification or the source is unclear, do not guess. Flag it for confirmation instead of deleting useful context.
5. Remove repetition only where the meaning is genuinely repeated. Do not remove necessary context just to make the paragraph shorter.
6. Rebuild awkward sentences naturally instead of replacing words one by one.
7. Use professional, natural Indian English suitable for CSR reports, Coffee Table Books, annual reports and donor publications.
8. Avoid exaggerated, promotional, dramatic or template style conclusions.
9. For beneficiary stories, preserve the actual sequence: previous situation, reason for joining, training or support, what happened afterwards, present livelihood or enterprise situation and practical change, only when supported by the source.
10. If a beneficiary already had a skill, say the programme strengthened, improved or commercialised it. Do not claim the programme created the skill.
11. Maintain consistent terminology, headings, abbreviations, trade names, programme names and location presentation.
12. Identify unsupported causal claims and mark them for confirmation.
13. Do not use hyphens, en dashes or em dashes in newly written prose.
14. STRICT LAYOUT RULE: For designed publications, keep the revised paragraph close to the original character count and overall length. Aim to remain within about 92 to 108 percent of the original character count unless there is clear duplication, unsupported content or a factual problem. Do not heavily shorten a paragraph merely to make it cleaner. If a major reduction seems necessary, retain the supported substance and flag the issue for confirmation.
15. Do not silently correct factual conflicts. Flag them.
16. Do not add new facts from general knowledge. Work only from the supplied text and supplied context.
17. Return exactly one JSON object. Do not return more than one JSON object. Do not add commentary before or after the JSON.

Return this shape only:
{
  "items": [
    {
      "id": 1,
      "revised": "revised paragraph",
      "reason": "short explanation of the main editorial change",
      "flags": ["Please confirm ..."]
    }
  ]
}
The id must match the paragraph id supplied by the user. Revise only the paragraphs listed under target_paragraphs. Use context_before and context_after only to understand flow.
"""

REFINE_PROMPT = r"""
You are revising one paragraph after the user has responded to an editorial confirmation question.
Use only the original paragraph, the current revision, the confirmation question, the user's response and the supplied surrounding context.

Interpret the user's response as follows:
1. YES means the flagged statement is confirmed and may be retained naturally.
2. NO means the flagged statement is not confirmed. Remove it or rewrite the paragraph so it does not make that unsupported claim.
3. COMMENT means use the user's written clarification as the verified basis for the rewrite.
4. Preserve all other verified facts, figures, names and locations exactly.
5. Treat place names as protected details. Preserve relevant state, district, taluk, block, town, village, Gram Panchayat, city, programme, training, market and other geographic references from India or any other country when they are needed for representation or context.
6. Respect the surrounding chapter or subject flow. The revised paragraph must continue naturally from the previous paragraph and connect logically with the following paragraph. Do not remove information needed for that continuity.
7. STRICT LAYOUT RULE: Keep the final paragraph close to the original character count, preferably within about 92 to 108 percent, unless the user's correction genuinely requires a larger change.
8. Do not add facts that are not in the original paragraph, surrounding context or the user's clarification.
9. Do not use hyphens, en dashes or em dashes in newly written prose.
10. Return exactly one JSON object with this shape only:
{
  "revised": "final revised paragraph",
  "reason": "short note explaining how the user's response was applied",
  "flags": []
}
If another genuinely unresolved factual issue remains, include it in flags. Otherwise return an empty flags list.
"""

ALLOWED_MODELS = {
    "gemini-3.1-flash-lite": "Gemini 3.1 Flash Lite",
    "gemini-3.8-flash": "Gemini 3.8 Flash",
}


def get_gemini_key() -> str:
    key = os.getenv("GEMINI_API_KEY")
    if not key:
        raise HTTPException(status_code=500, detail="GEMINI_API_KEY is not configured on the server. Create a free Gemini API key in Google AI Studio and add it to Render Environment.")
    return key


def extract_text(filename: str, data: bytes) -> List[str]:
    ext = Path(filename).suffix.lower()
    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        return [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        paras: List[str] = []
        for page in reader.pages:
            text = page.extract_text() or ""
            for block in re.split(r"\n\s*\n|\n(?=[A-Z0-9])", text):
                block = re.sub(r"\s+", " ", block).strip()
                if block:
                    paras.append(block)
        return paras
    if ext in {".txt", ".md"}:
        text = data.decode("utf-8", errors="ignore")
        return [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    raise HTTPException(status_code=400, detail="Supported files: DOCX, PDF, TXT and MD.")


def text_to_paragraphs(text: str) -> List[str]:
    text = (text or "").strip()
    if not text:
        return []
    parts = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    return parts if parts else [text]


def chunks(items: List[Dict[str, Any]], max_chars: int = 6000):
    batch, size = [], 0
    for item in items:
        s = len(item["text"])
        if batch and size + s > max_chars:
            yield batch
            batch, size = [], 0
        batch.append(item)
        size += s
    if batch:
        yield batch


def _strip_fences(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.I)
    text = re.sub(r"\s*```$", "", text)
    return text.strip()


def extract_json(text: str) -> Dict[str, Any]:
    text = _strip_fences(text)
    decoder = json.JSONDecoder()
    objects = []
    pos = 0
    while pos < len(text):
        while pos < len(text) and text[pos] in " \t\r\n,;":
            pos += 1
        if pos >= len(text):
            break
        if text[pos] not in "[{":
            nxt = min([x for x in (text.find("{", pos), text.find("[", pos)) if x != -1], default=-1)
            if nxt == -1:
                break
            pos = nxt
        try:
            obj, end = decoder.raw_decode(text, pos)
            objects.append(obj)
            pos = end
        except json.JSONDecodeError:
            pos += 1
    if not objects:
        raise ValueError("Gemini did not return readable JSON. Please try again.")
    merged_items = []
    for obj in objects:
        if isinstance(obj, dict) and isinstance(obj.get("items"), list):
            merged_items.extend(obj["items"])
        elif isinstance(obj, list):
            merged_items.extend(x for x in obj if isinstance(x, dict))
    if merged_items:
        return {"items": merged_items}
    if isinstance(objects[0], dict):
        return objects[0]
    raise ValueError("Gemini JSON did not contain the expected object.")


def gemini_generate(model: str, payload: str, system_prompt: str = SYSTEM_PROMPT) -> str:
    key = get_gemini_key()
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
    body = {
        "systemInstruction": {"parts": [{"text": system_prompt}]},
        "contents": [{"role": "user", "parts": [{"text": payload}]}],
        "generationConfig": {"temperature": 0.1, "responseMimeType": "application/json", "maxOutputTokens": 32768},
    }
    try:
        with httpx.Client(timeout=150.0) as client:
            response = client.post(url, headers={"x-goog-api-key": key, "Content-Type": "application/json"}, json=body)
    except httpx.HTTPError as exc:
        raise RuntimeError(f"Could not reach Gemini API: {exc}") from exc
    if response.status_code >= 400:
        try:
            err = response.json().get("error", {})
            message = err.get("message") or response.text
        except Exception:
            message = response.text
        if response.status_code == 429:
            raise RuntimeError("Gemini free tier rate limit reached. Please wait and try again later.")
        if response.status_code in {401, 403}:
            raise RuntimeError("Gemini API key is invalid or does not have access to this model.")
        raise RuntimeError(f"Gemini API error {response.status_code}: {message}")
    data = response.json()
    candidates = data.get("candidates") or []
    if not candidates:
        raise RuntimeError(f"Gemini returned no editorial response. {data.get('promptFeedback') or {}}")
    parts = candidates[0].get("content", {}).get("parts", [])
    text = "".join(str(part.get("text", "")) for part in parts if part.get("text"))
    if not text.strip():
        raise RuntimeError("Gemini returned an empty editorial response.")
    return text


def review_paragraphs(paragraphs: List[str], model: str) -> List[Dict[str, Any]]:
    source_items = [{"id": i + 1, "text": p} for i, p in enumerate(paragraphs)]
    reviewed: Dict[int, Dict[str, Any]] = {}

    for batch in chunks(source_items):
        first_id = batch[0]["id"]
        last_id = batch[-1]["id"]
        context_before = [
            {"id": i + 1, "text": paragraphs[i]}
            for i in range(max(0, first_id - 4), first_id - 1)
        ]
        context_after = [
            {"id": i + 1, "text": paragraphs[i]}
            for i in range(last_id, min(len(paragraphs), last_id + 3))
        ]
        payload = json.dumps({
            "instruction": "Revise only target_paragraphs. Use context_before and context_after to preserve chapter or subject continuity.",
            "context_before": context_before,
            "target_paragraphs": batch,
            "context_after": context_after,
        }, ensure_ascii=False)
        data = extract_json(gemini_generate(model, payload))
        for item in data.get("items", []):
            if not isinstance(item, dict):
                continue
            try:
                idx = int(item.get("id"))
            except Exception:
                continue
            if idx < 1 or idx > len(paragraphs):
                continue
            flags = item.get("flags", []) or []
            if not isinstance(flags, list):
                flags = [str(flags)]
            reviewed[idx] = {
                "id": idx,
                "original": paragraphs[idx - 1],
                "revised": str(item.get("revised", paragraphs[idx - 1])).strip(),
                "reason": str(item.get("reason", "")).strip(),
                "flags": [str(flag) for flag in flags],
            }

    return [reviewed.get(i, {
        "id": i,
        "original": original,
        "revised": original,
        "reason": "No change returned by the model.",
        "flags": ["Please review manually because no model revision was returned."],
    }) for i, original in enumerate(paragraphs, start=1)]


def add_diff_paragraph(doc: Document, original: str, revised: str):
    p = doc.add_paragraph()
    sm = difflib.SequenceMatcher(None, original.split(), revised.split())
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            r = p.add_run(" ".join(original.split()[i1:i2]) + " ")
            r.font.color.rgb = RGBColor(0, 0, 0)
        elif tag in {"delete", "replace"}:
            deleted = " ".join(original.split()[i1:i2])
            if deleted:
                r = p.add_run(deleted + " ")
                r.font.color.rgb = RGBColor(192, 0, 0)
                r.font.strike = True
            if tag == "replace":
                added = " ".join(revised.split()[j1:j2])
                if added:
                    r = p.add_run(added + " ")
                    r.font.color.rgb = RGBColor(0, 102, 204)
        elif tag == "insert":
            added = " ".join(revised.split()[j1:j2])
            if added:
                r = p.add_run(added + " ")
                r.font.color.rgb = RGBColor(0, 102, 204)
    return p


def build_review_doc(results: List[Dict[str, Any]], title: str) -> bytes:
    doc = Document()
    doc.add_heading("Editorial Review", level=1)
    doc.add_paragraph(title)
    legend = doc.add_paragraph()
    r = legend.add_run("Red strikethrough")
    r.font.color.rgb = RGBColor(192, 0, 0)
    r.font.strike = True
    legend.add_run(" means deletion. ")
    r = legend.add_run("Blue text")
    r.font.color.rgb = RGBColor(0, 102, 204)
    legend.add_run(" means addition or replacement.")
    for item in results:
        doc.add_heading(f"Paragraph {item['id']}", level=2)
        add_diff_paragraph(doc, item["original"], item["revised"])
        if item["reason"]:
            p = doc.add_paragraph()
            p.add_run("Editorial note: ").bold = True
            p.add_run(item["reason"])
        for flag in item["flags"]:
            p = doc.add_paragraph()
            rr = p.add_run("Please confirm: ")
            rr.bold = True
            rr.font.color.rgb = RGBColor(192, 0, 0)
            p.add_run(str(flag))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def build_clean_doc(results: List[Dict[str, Any]], title: str) -> bytes:
    doc = Document()
    doc.add_heading(title, level=1)
    for item in results:
        doc.add_paragraph(item["revised"])
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    return templates.TemplateResponse("index.html", {"request": request})


@app.get("/health")
def health():
    return {"status": "ok", "provider": "gemini", "gemini_key_configured": bool(os.getenv("GEMINI_API_KEY"))}


@app.post("/review")
async def review(
    file: Optional[UploadFile] = File(None),
    text: str = Form(""),
    model: str = Form("gemini-3.1-flash-lite"),
):
    if model not in ALLOWED_MODELS:
        raise HTTPException(status_code=400, detail="Unsupported model selection.")

    typed_text = (text or "").strip()
    if typed_text:
        if len(typed_text) > 120000:
            raise HTTPException(status_code=400, detail="Pasted text is too long. Please review it in smaller sections.")
        paragraphs = text_to_paragraphs(typed_text)
        source_name = "Pasted Text"
    elif file and file.filename:
        data = await file.read()
        if len(data) > 15 * 1024 * 1024:
            raise HTTPException(status_code=400, detail="File is too large. Maximum size is 15 MB.")
        paragraphs = extract_text(file.filename, data)
        source_name = file.filename
    else:
        raise HTTPException(status_code=400, detail="Upload a document or paste/type text to review.")

    if not paragraphs:
        raise HTTPException(status_code=400, detail="No readable text was found.")
    try:
        results = review_paragraphs(paragraphs, model)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Editorial review could not be completed: {exc}") from exc
    return {"filename": source_name, "count": len(results), "items": results, "provider": "Gemini free tier"}


@app.post("/refine")
async def refine(
    paragraph_id: int = Form(...),
    original: str = Form(...),
    revised: str = Form(...),
    flag: str = Form(""),
    decision: str = Form(...),
    comment: str = Form(""),
    context_before: str = Form(""),
    context_after: str = Form(""),
    model: str = Form("gemini-3.1-flash-lite"),
):
    if model not in ALLOWED_MODELS:
        raise HTTPException(status_code=400, detail="Unsupported model selection.")
    decision = decision.strip().upper()
    if decision not in {"YES", "NO", "COMMENT"}:
        raise HTTPException(status_code=400, detail="Decision must be YES, NO or COMMENT.")
    if decision == "COMMENT" and not comment.strip():
        raise HTTPException(status_code=400, detail="Please write your clarification before rewriting.")

    payload = json.dumps({
        "paragraph_id": paragraph_id,
        "original": original,
        "current_revision": revised,
        "confirmation_question": flag,
        "user_response_type": decision,
        "user_comment": comment.strip(),
        "context_before": context_before.strip(),
        "context_after": context_after.strip(),
        "original_character_count": len(original),
    }, ensure_ascii=False)
    try:
        data = extract_json(gemini_generate(model, payload, REFINE_PROMPT))
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Paragraph rewrite could not be completed: {exc}") from exc

    flags = data.get("flags", []) or []
    if not isinstance(flags, list):
        flags = [str(flags)]
    return {
        "id": paragraph_id,
        "original": original,
        "revised": str(data.get("revised", revised)).strip(),
        "reason": str(data.get("reason", "Updated based on your confirmation.")).strip(),
        "flags": [str(x) for x in flags],
    }


@app.post("/download/reviewed")
async def download_reviewed(payload: str = Form(...), filename: str = Form("document.docx")):
    results = json.loads(payload)
    data = build_review_doc(results, filename)
    stem = Path(filename).stem if filename else "pasted_text"
    return StreamingResponse(io.BytesIO(data), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f'attachment; filename="{stem}_editorial_review.docx"'})


@app.post("/download/clean")
async def download_clean(payload: str = Form(...), filename: str = Form("document.docx")):
    results = json.loads(payload)
    stem = Path(filename).stem if filename else "pasted_text"
    data = build_clean_doc(results, f"{stem} Revised Copy")
    return StreamingResponse(io.BytesIO(data), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f'attachment; filename="{stem}_clean_revised.docx"'})