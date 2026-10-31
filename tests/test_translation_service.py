import io
import json

import pytest
from docx import Document
from fastapi.testclient import TestClient

import app as review_app
from translation_service import (
    build_translation_docx,
    count_source_words,
    extract_translation_segments,
    translation_needs_review,
)


def sample_docx():
    doc = Document()
    doc.add_heading("हिन्दी शीर्षक", level=1)
    doc.add_paragraph("नमस्ते दुनिया")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "कैसे हो"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_hindi_unicode_word_count():
    assert count_source_words("नमस्ते दुनिया 2026!") == 3
    assert count_source_words("માતૃભાષા शिक्षा") == 2
    assert count_source_words("  ") == 0


def test_docx_segments_keep_heading_and_table_order():
    passages = extract_translation_segments("original.docx", sample_docx())
    assert [p["original"] for p in passages] == [
        "हिन्दी शीर्षक", "नमस्ते दुनिया", "कैसे हो"
    ]


def test_docx_export_preserves_styles_and_tables():
    source = sample_docx()
    items = extract_translation_segments("source.docx", source)
    words = ["ગુજરાતી શીર્ષક", "નમસ્તે દુનિયા", "કેમ છો"]
    for item, target in zip(items, words):
        item["translated"] = target
    new_bytes = build_translation_docx(items, "source.docx", source)
    translated = Document(io.BytesIO(new_bytes))
    assert translated.paragraphs[0].style.name == "Heading 1"
    assert translated.paragraphs[0].text == words[0]
    assert translated.paragraphs[1].text == words[1]
    assert translated.tables[0].cell(0, 0).text == words[2]


def test_docx_export_rejects_source_mismatch():
    items = extract_translation_segments("source.docx", sample_docx())
    for item in items:
        item["translated"] = "ગુજરાતી"
    items[0]["original"] = "changed"
    with pytest.raises(ValueError, match="source content has changed"):
        build_translation_docx(items, "source.docx", sample_docx())


def test_untranslated_hindi_flagged():
    assert translation_needs_review("नमस्ते", "नमस्ते")
    assert translation_needs_review("नमस्ते", "નમસ્તે") == []


def test_api_estimate_translate_and_verified_download(monkeypatch):
    def fake_gemini(model, payload, system_prompt):
        data = json.loads(payload)
        return json.dumps({
            "items": [
                {"id": item["id"], "translated": "ગુજરાતી અનુવાદ", "flags": []}
                for item in data["passages"]
            ]
        }, ensure_ascii=False)

    monkeypatch.setattr(review_app, "gemini_generate", fake_gemini)
    client = TestClient(review_app.app)
    source = sample_docx()
    files = {"file": ("source.docx", source, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    generated = client.post("/translation/translate", files=files)
    assert generated.status_code == 200, generated.text
    body = generated.json()
    assert len(body["items"]) == 3
    assert all(item["reviewed"] is False for item in body["items"])
    assert all(item["translated"] == "ગુજરાતી અનુવાદ" for item in body["items"])

    data = {
        "filename": "source.docx",
        "source_sha256": body["source_sha256"],
        "stage": "final",
        "payload": json.dumps(body["items"], ensure_ascii=False),
    }
    blocked = client.post("/translation/download", data=data, files=files)
    assert blocked.status_code == 400
    assert "verified" in blocked.json()["detail"]

    for item in body["items"]:
        item["reviewed"] = True
    data["payload"] = json.dumps(body["items"], ensure_ascii=False)
    approved = client.post("/translation/download", data=data, files=files)
    assert approved.status_code == 200, approved.text
    assert approved.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'Gujarati_final.docx' in approved.headers["content-disposition"]
    doc = Document(io.BytesIO(approved.content))
    assert doc.tables[0].cell(0, 0).text == "ગુજરાતી અનુવાદ"


def test_revision_round_limit():
    client = TestClient(review_app.app)
    response = client.post("/translation/revise", data={
        "original": "हिन्दी",
        "translated": "ગુજરાતી",
        "feedback": "Please correct terminology",
        "revision_round": "5",
    })
    assert response.status_code == 400


def test_cannot_download_without_translation():
    client = TestClient(review_app.app)
    response = client.post("/translation/download", data={
        "payload": json.dumps([{"id": 1, "original": "नमस्ते", "translated": ""}]),
        "filename": "text.txt",
    })
    assert response.status_code == 400
