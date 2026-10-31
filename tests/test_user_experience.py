"""Regression coverage for usability and document accuracy improvements."""
import io
import json
import re
import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

import app as review_app
from translation_service import build_translation_docx, translation_needs_review


def test_page_has_user_friendly_translation_tools():
    result = TestClient(review_app.app).get("/")
    assert result.status_code == 200
    html = result.text
    for control in [
        'id="translationAudience"', 'id="translationGlossary"',
        'id="translationSaveProject"', 'id="translationLoadProject"',
        'id="translationCopyAll"', 'id="translationTextDownload"',
        'id="translationOriginalAgain"', 'id="translationNewProject"',
    ]:
        assert control in html
    assert "n*m>180000" in html


def test_inline_javascript_parses_when_node_available():
    node = shutil.which("node")
    if node is None:
        pytest.skip("Node.js unavailable on runner")
    markup = (Path(__file__).resolve().parents[1] / "templates" / "index.html").read_text()
    match = re.search(r"<script>(.*?)</script>", markup, re.DOTALL)
    assert match, "Missing application JavaScript"
    checked = subprocess.run([node, "--check"], input=match.group(1),
                             text=True, capture_output=True)
    assert checked.returncode == 0, checked.stderr


def test_institutional_style_and_terms_reach_translation_and_revision(monkeypatch):
    seen = []

    def fake_gemini(model, payload, system_prompt):
        request = json.loads(payload)
        seen.append(request)
        if "passages" in request:
            return json.dumps({"items": [
                {"id": item["id"], "translated": "नमस्ते समुदाय 2026", "flags": []}
                for item in request["passages"]
            ]}, ensure_ascii=False)
        return json.dumps({"translated": "नमस्ते समुदाय 2026", "flags": []}, ensure_ascii=False)

    monkeypatch.setattr(review_app, "gemini_generate", fake_gemini)
    client = TestClient(review_app.app)
    options = {
        "source_language": "en", "target_language": "hi",
        "audience": "ngo", "glossary": "community = समुदाय\ntraining = प्रशिक्षण"
    }
    result = client.post("/translation/translate", data={**options, "text": "Hello community 2026"})
    assert result.status_code == 200, result.text
    assert result.json()["writing_style"] == "ngo"
    assert seen[0]["source_language"] == "English"
    assert seen[0]["target_language"] == "Hindi"
    assert "NGO" in seen[0]["writing_style"]
    assert seen[0]["preferred_terminology"] == [
        {"source": "community", "target": "समुदाय"},
        {"source": "training", "target": "प्रशिक्षण"},
    ]
    revised = client.post("/translation/revise", data={
        **options, "original": "Hello community 2026",
        "translated": "नमस्ते समुदाय 2026", "revision_round": 1,
        "feedback": "Keep the approved programme terms",
    })
    assert revised.status_code == 200, revised.text
    assert seen[1]["preferred_terminology"] == seen[0]["preferred_terminology"]


@pytest.mark.parametrize("audience,glossary", [
    ("unverified", ""), ("ngo", "incomplete term"),
    ("formal", "a = "), ("education", "x = y\n" * 31),
])
def test_reject_invalid_translation_settings(audience, glossary):
    client = TestClient(review_app.app)
    result = client.post("/translation/translate", data={
        "source_language": "en", "target_language": "hi",
        "text": "Hello world", "audience": audience, "glossary": glossary,
    })
    assert result.status_code == 400, result.text


def test_amount_and_number_mismatch_flagged():
    assert any("Numbers or amounts" in flag for flag in translation_needs_review(
        "Funding 12,500 in 2026", "2026 में 12,000 का अनुदान", "hi"
    ))
    assert any("Numbers or amounts" in flag for flag in translation_needs_review(
        "Funding 12,500 in 2026", "2026 में 12,500 का अनुदान", "hi"
    ))  # Same figures but reversed order still requires a human check
    assert not any("Numbers or amounts" in flag for flag in translation_needs_review(
        "Funding 12,500 in 2026", "12,500 का अनुदान 2026 में", "hi"
    ))


def test_right_to_left_docx_output():
    items = [{"id": 1, "original": "Welcome", "translated": "خوش آمدید"}]
    content = build_translation_docx(items, "typed.txt", target_language="ur")
    doc = Document(io.BytesIO(content))
    assert doc.paragraphs[0].text == "خوش آمدید"
    assert doc.paragraphs[0]._p.xpath("./w:pPr/w:bidi")


def test_download_endpoints_reject_invalid_review_data():
    client = TestClient(review_app.app)
    for endpoint in ("/download/clean", "/download/reviewed"):
        for payload in ("not json", "{}", '[{"revised": "x"}]'):
            result = client.post(endpoint, data={"payload": payload, "filename": "report.docx"})
            assert result.status_code == 400


def test_translation_download_rejects_empty_project():
    result = TestClient(review_app.app).post("/translation/download", data={"payload": "[]"})
    assert result.status_code == 400
